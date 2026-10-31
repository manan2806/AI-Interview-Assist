import csv, io, logging
from datetime import datetime, timedelta
from flask import Blueprint, jsonify, request, make_response

from database import db
from utils.auth_utils import admin_required
from flask import send_file

from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from docx import Document

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

reports_bp = Blueprint("reports", __name__)


def get_nested_value(document, field):
    """Get values from nested MongoDB documents using dot notation."""
    value = document

    for key in field.split("."):
        if not isinstance(value, dict):
            return ""

        value = value.get(key)

        if value is None:
            return ""

    return value


# ==========================================
# REPORT CONFIGURATION
# ==========================================
REPORT_CONFIG = {
    "users": {
        "collection": "users",
        "fields": [
            "user_id",
            "name",
            "email",
            "target_role",
            "experience_level",
            "skills",
        ],
        "date_field": "created_at",
        "exclude_admins": True,
    },
    "interviews": {
        "collection": "interviews",
        "fields": [
            "interview_code",
            "user_id",
            "job_role",
            "experience_level",
            "interview_type",
            "difficulty",
            "status",
            "overall_result.overall_score",
            "overall_result.performance_level",
            "created_at",
        ],
        "date_field": "created_at",
        "exclude_admins": False,
    },
    "problem-reports": {
        "collection": "problem_reports",
        "fields": [
            "report_id",
            "user_id",
            "user_name",
            "user_email",
            "subject",
            "problem_type",
            "description",
            "status",
            "created_at",
        ],
        "date_field": "created_at",
        "exclude_admins": False,
    },
    "reviews": {
        "collection": "reviews",
        "fields": [
            "review_id",
            "user_id",
            "name",
            "rating",
            "comment",
            "status",
            "created_at",
            "updated_at",
            "moderated_by",
        ],
        "date_field": "created_at",
        "exclude_admins": False,
    },
}


# ==========================================
# CSV VALUE HELPER
# ==========================================
def report_value(value):
    """Convert MongoDB values into CSV-friendly text."""
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, (dict, list, tuple)):
        value = str(value)
    else:
        value = str(value)

    # Reduce spreadsheet formula-injection risk.
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value

    return value


# ==========================================
# REPORT FILE GENERATION HELPER
# ==========================================
def generate_report_file(report_type, fields, rows, export_format):
    filename = f"{report_type}-report"

    # CSV
    if export_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(fields)
        writer.writerows(rows)

        buffer = io.BytesIO()
        buffer.write(("\ufeff" + output.getvalue()).encode("utf-8"))
        buffer.seek(0)

        return buffer, "text/csv", f"{filename}.csv"

    # PDF
    elif export_format == "pdf":
        buffer = io.BytesIO()

        page_width, page_height = landscape(A3)

        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A3),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25,
        )

        styles = getSampleStyleSheet()

        header_style = ParagraphStyle(
            "ReportHeader",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7,
            leading=9,
            textColor=colors.white,
            alignment=1,
        )

        cell_style = ParagraphStyle(
            "ReportCell",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            wordWrap="CJK",
        )

        elements = [
            Paragraph(
                f"{report_type.replace('-', ' ').title()} Report",
                styles["Title"],
            ),
            Spacer(1, 15),
        ]

        header_labels = {
            "interview_code": "Interview Code",
            "user_id": "User ID",
            "job_role": "Job Role",
            "experience_level": "Experience Level",
            "interview_type": "Interview Type",
            "difficulty": "Difficulty",
            "status": "Status",
            "overall_result.overall_score": "Overall Score",
            "overall_result.performance_level": "Performance Level",
            "created_at": "Created At",
        }

        table_data = [
            [
                Paragraph(
                    header_labels.get(field, field.replace("_", " ").title()),
                    header_style,
                )
                for field in fields
            ]
        ]

        for row in rows:
            table_data.append(
                [
                    Paragraph(
                        str(value if value is not None else "")
                        .replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;"),
                        cell_style,
                    )
                    for value in row
                ]
            )

        available_width = page_width - 50

        if report_type == "interviews":
            weights = [
                1.1,  # Interview Code
                0.9,  # User ID
                1.5,  # Job Role
                1.2,  # Experience Level
                1.1,  # Interview Type
                0.8,  # Difficulty
                0.9,  # Status
                0.9,  # Overall Score
                1.3,  # Performance Level
                1.5,  # Created At
            ]
        else:
            weights = [1] * len(fields)

        total_weight = sum(weights)

        col_widths = [available_width * weight / total_weight for weight in weights]

        table = Table(
            table_data,
            colWidths=col_widths,
            repeatRows=1,
            splitByRow=1,
        )

        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5146d8")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor("#f2f3f8"),
                        ],
                    ),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )

        elements.append(table)
        document.build(elements)

        buffer.seek(0)

        return (
            buffer,
            "application/pdf",
            f"{filename}.pdf",
        )

    # Word
    elif export_format == "docx":
        document = Document()
        document.add_heading(
            f"{report_type.replace('-', ' ').title()} Report",
            level=1,
        )

        table = document.add_table(rows=1, cols=len(fields))
        table.style = "Table Grid"

        for index, field in enumerate(fields):
            table.rows[0].cells[index].text = str(field)

        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)

        return (
            buffer,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{filename}.docx",
        )

    # Excel
    elif export_format == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        worksheet.append(fields)

        for row in rows:
            worksheet.append(row)

        header_fill = PatternFill(
            start_color="5146D8",
            end_color="5146D8",
            fill_type="solid",
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for column in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            worksheet.column_dimensions[column[0].column_letter].width = min(
                max_length + 3, 40
            )

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        return (
            buffer,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{filename}.xlsx",
        )

    else:
        raise ValueError("Unsupported export format.")


# ==========================================
# ADMIN REPORT EXPORT API
# ==========================================
@reports_bp.route("/api/admin/reports/<report_type>/export", methods=["GET"])
@admin_required
def export_admin_report(current_user, report_type):

    try:
        if db is None:
            return (
                jsonify({"success": False, "message": "Database is not connected."}),
                500,
            )

        # REPORT TYPE VALIDATION
        config = REPORT_CONFIG.get(report_type)

        if not config:
            return (
                jsonify(
                    {
                        "success": False,
                        "message": "Unsupported report type.",
                        "allowed_types": list(REPORT_CONFIG.keys()),
                    }
                ),
                400,
            )

        # EXPORT FORMAT VALIDATION
        export_format = request.args.get("format", "csv").strip().lower()

        allowed_formats = ["csv", "pdf", "docx", "xlsx"]

        if export_format not in allowed_formats:
            return (
                jsonify(
                    {
                        "success": False,
                        "message": "Unsupported export format.",
                        "allowed_formats": allowed_formats,
                    }
                ),
                400,
            )

        # BASE QUERY
        query = {}

        if config["exclude_admins"]:
            query["role"] = {"$ne": "admin"}

        # DATE FILTER
        start_date = request.args.get("start_date", "").strip()

        end_date = request.args.get("end_date", "").strip()

        date_query = {}

        if start_date:
            try:
                start = datetime.strptime(start_date, "%Y-%m-%d")

                date_query["$gte"] = start

            except ValueError:
                return (
                    jsonify(
                        {
                            "success": False,
                            "message": "Invalid start_date. Use YYYY-MM-DD.",
                        }
                    ),
                    400,
                )

        if end_date:
            try:
                end = datetime.strptime(end_date, "%Y-%m-%d")

                date_query["$lt"] = end + timedelta(days=1)

            except ValueError:
                return (
                    jsonify(
                        {
                            "success": False,
                            "message": "Invalid end_date. Use YYYY-MM-DD.",
                        }
                    ),
                    400,
                )

        if start_date and end_date and start_date > end_date:
            return (
                jsonify(
                    {
                        "success": False,
                        "message": "Start date cannot be after end date.",
                    }
                ),
                400,
            )

        if date_query:
            query[config["date_field"]] = date_query

        # FETCH REPORT DATA
        collection = db[config["collection"]]
        fields = config["fields"]

        cursor = collection.find(query, {field: 1 for field in fields}).sort(
            config["date_field"], -1
        )

        rows = []

        for document in cursor:
            row = [report_value(get_nested_value(document, field)) for field in fields]

            rows.append(row)

        # GENERATE REPORT FILE
        buffer, mimetype, download_name = generate_report_file(
            report_type, fields, rows, export_format
        )

        # SEND DOWNLOAD RESPONSE
        response = send_file(
            buffer, mimetype=mimetype, as_attachment=True, download_name=download_name
        )

        response.headers["X-Exported-Count"] = str(len(rows))

        return response

    except Exception:
        logging.exception("Admin report export error")

        return jsonify({"success": False, "message": "Unable to export report."}), 500
